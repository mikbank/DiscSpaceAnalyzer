import os
import sys
from concurrent.futures import FIRST_COMPLETED, ProcessPoolExecutor, wait
from datetime import datetime

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

load_dotenv()

SCAN_ROOT = os.getenv("SCAN_ROOT", "C:\\")
SIZE_THRESHOLD_GB = float(os.getenv("SIZE_THRESHOLD_GB", "1"))
OUTPUT_FILE = os.getenv("OUTPUT_FILE", "large_folders_report.xlsx")
WORKERS = int(os.getenv("WORKERS", "8"))

GB = 1024 ** 3
THRESHOLD_BYTES = SIZE_THRESHOLD_GB * GB

# If a subtree turns out to hold more than this many folders, a task bails
# out of scanning it and instead reports its immediate children as new,
# smaller tasks. This is what a raw subdirectory-count check can't get
# right: a folder can have just one or two children (e.g. C:\Users with a
# handful of profiles) where a single child is enormous. Reacting to the
# actual measured size, recursively, keeps splitting wherever the real
# size is -- however deep that turns out to be -- instead of leaving one
# worker stuck scanning a disproportionately large folder alone while the
# others sit idle.
SPLIT_CAP_FOLDERS = 2000

# When a directory needs splitting, its children are handed out in
# batches of this size rather than one task per child -- otherwise a
# directory with tens of thousands of direct children (again, WinSxS is
# a real example) would explode into as many process-pool tasks, and the
# per-task overhead alone would dominate the runtime.
BATCH_SIZE = 100


def _list_dir(path):
    """One scandir pass over `path`: bytes used by files directly inside
    it, and its real (non-symlink) subdirectories.

    Size comes from the DirEntry's own stat() call, which Windows
    populates from the same directory-enumeration data scandir already
    fetched -- no extra per-file syscall, unlike os.path.getsize().
    Errors (permission denied, locked files, etc.) are swallowed so one
    bad folder just contributes 0 instead of aborting the scan.

    Symlinked directories are skipped (not recursed into) to avoid
    cycles, matching os.walk's default followlinks=False behavior. Note
    this checks is_symlink() specifically, not the broader "reparse
    point" attribute -- cloud-sync folders (OneDrive, etc.) also carry a
    reparse point flag while being ordinary real directories, and must
    still be scanned.
    """
    own_total = 0
    subdirs = []
    try:
        with os.scandir(path) as it:
            for entry in it:
                try:
                    if entry.is_dir(follow_symlinks=False):
                        if entry.is_symlink():
                            continue
                        subdirs.append(entry.path)
                    else:
                        own_total += entry.stat(follow_symlinks=False).st_size
                except OSError:
                    continue
    except OSError:
        pass
    return own_total, subdirs


def _scan_capped(root, split_cap):
    """Compute the size of every folder under (and including) root, the
    same way as a plain bottom-up scan -- except it gives up as soon as
    it has visited more than `split_cap` folders, returning `root`'s own
    file total and its direct children instead of finishing the job.

    Bottom-up via an explicit stack rather than recursion, so a very deep
    tree can't hit Python's recursion limit. `root` is always the first
    item processed (it's the only item on the stack to start), so its own
    total/children are available even if the cap trips on the very next
    directory.
    """
    own_totals = {}
    children = {}
    order = []
    stack = [root]

    while stack:
        if len(order) >= split_cap:
            return None, own_totals[root], children[root]
        current = stack.pop()
        order.append(current)
        own_total, subdirs = _list_dir(current)
        own_totals[current] = own_total
        children[current] = subdirs
        stack.extend(subdirs)

    sizes = {}
    for path in reversed(order):
        total = own_totals[path]
        for sub in children[path]:
            total += sizes.get(sub, 0)
        sizes[path] = total

    return sizes, None, None


def scan_subtree(root):
    """Uncapped version of _scan_capped, for direct/testing use."""
    sizes, _, _ = _scan_capped(root, split_cap=float("inf"))
    return sizes


def _chunked(items, size):
    for i in range(0, len(items), size):
        yield items[i:i + size]


def _scan_task(paths, threshold_bytes, split_cap):
    """Runs in a worker process: scan a batch of independent subtree
    roots. Filtering to just the large folders on success (instead of
    shipping every folder's size back) keeps inter-process traffic small
    even for a subtree with hundreds of thousands of folders.

    Batching matters because a directory can have an enormous number of
    direct children that are each individually tiny (Windows' WinSxS
    component store is a real example: tens of thousands of near-empty
    folders under one parent). Submitting one process-pool task per child
    in a case like that would itself become the bottleneck. Any path that
    turns out too big to finish here is reported back separately, so the
    caller can re-batch just its children into new tasks.
    """
    done = []
    needs_split = []
    for path in paths:
        sizes, own_total, subdirs = _scan_capped(path, split_cap)
        if sizes is None:
            needs_split.append((path, own_total, subdirs))
        else:
            large = [(p, s) for p, s in sizes.items() if s >= threshold_bytes]
            done.append((path, sizes.get(path, 0), large, len(sizes)))
    return done, needs_split


def write_excel(large_folders, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Large Folders"

    headers = ["Path", "Size (GB)", "Size (Bytes)"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).font = Font(bold=True)

    for path, size in large_folders:
        ws.append([path, round(size / GB, 2), size])

    widths = [80, 14, 16]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(output_file)


def main():
    if not os.path.isdir(SCAN_ROOT):
        print(f"SCAN_ROOT does not exist or is not a folder: {SCAN_ROOT}")
        sys.exit(1)

    print(f"Scanning {SCAN_ROOT} for folders >= {SIZE_THRESHOLD_GB} GB using {WORKERS} workers...")
    started = datetime.now()

    # Directories that got split: {dir: (own_total, [child dirs])}, needed
    # to re-aggregate their sizes once their children's results are in.
    intermediate = {}
    sizes_by_path = {}
    large_folders = []
    total_dir_count = 0

    root_own_total, root_subdirs = _list_dir(SCAN_ROOT)
    if root_subdirs:
        intermediate[SCAN_ROOT] = (root_own_total, root_subdirs)
        initial_paths = root_subdirs
    else:
        initial_paths = [SCAN_ROOT]

    with ProcessPoolExecutor(max_workers=WORKERS) as executor:
        pending = {
            executor.submit(_scan_task, batch, THRESHOLD_BYTES, SPLIT_CAP_FOLDERS)
            for batch in _chunked(initial_paths, BATCH_SIZE)
        }
        completed = 0
        while pending:
            done, pending = wait(pending, return_when=FIRST_COMPLETED)
            for future in done:
                results, needs_split = future.result()
                for path, total, large, dir_count in results:
                    sizes_by_path[path] = total
                    large_folders.extend(large)
                    total_dir_count += dir_count
                    completed += 1
                    print(f"  ...chunk {completed} done: {dir_count} folders, {total / GB:.2f} GB "
                          f"({path}) -- {len(pending)} batches in flight")
                if needs_split:
                    more_paths = []
                    for path, own_total, subdirs in needs_split:
                        intermediate[path] = (own_total, subdirs)
                        more_paths.extend(subdirs)
                    pending |= {
                        executor.submit(_scan_task, batch, THRESHOLD_BYTES, SPLIT_CAP_FOLDERS)
                        for batch in _chunked(more_paths, BATCH_SIZE)
                    }

    # Re-aggregate the directories that were split, in reverse (children-
    # before-parents) order -- a directory's entry is only added to
    # `intermediate` after its own scan task has run, and a child can only
    # be split *after* its parent was, so insertion order is already a
    # valid top-down order and reversing it gives bottom-up.
    for path in reversed(intermediate):
        own_total, subdirs = intermediate[path]
        total = own_total
        for sub in subdirs:
            total += sizes_by_path.get(sub, 0)
        sizes_by_path[path] = total
        total_dir_count += 1
        if total >= THRESHOLD_BYTES:
            large_folders.append((path, total))

    large_folders.sort(key=lambda item: item[1], reverse=True)
    write_excel(large_folders, OUTPUT_FILE)

    elapsed = datetime.now() - started
    print(f"Scanned {total_dir_count} folders in {elapsed}.")
    print(f"Found {len(large_folders)} folders >= {SIZE_THRESHOLD_GB} GB.")
    print(f"Report written to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
